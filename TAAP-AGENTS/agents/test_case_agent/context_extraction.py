"""Preprocessing for non-Jira supplemental context (HTML / audio / files).

Converts the raw, untrusted payload a client sends in `SupplementalContext`
(Base64 file/audio content or raw HTML) into plain, normalized text the Test
Case Generation Agent can safely read. Raw Base64 is never passed to the LLM:
every code path here returns a `SupplementalContext` with `content_base64`
and `raw_html` cleared, leaving only `extracted_text` / `transcript`.
"""

import base64
import binascii
import csv
import io
import json
import logging
import os
import re
from typing import Callable, Optional

import litellm
from bs4 import BeautifulSoup, Comment
from docx import Document
from pypdf import PdfReader

from .models import SupplementalContext

logger = logging.getLogger(__name__)

AUDIO_MODEL_NAME = os.getenv("TEST_CASE_AUDIO_MODEL", "openai/whisper-1")

# Decoded-content size limits. Configurable via env to match whatever upload
# limits the surrounding application already enforces at the edge.
MAX_FILE_BYTES = int(
    os.getenv("TEST_CASE_SUPPLEMENTAL_FILE_MAX_BYTES", str(15 * 1024 * 1024))
)
MAX_AUDIO_BYTES = int(
    os.getenv("TEST_CASE_SUPPLEMENTAL_AUDIO_MAX_BYTES", str(25 * 1024 * 1024))
)

SUPPORTED_AUDIO_MIME_TYPES = {
    "audio/mpeg",
    "audio/mp3",
    "audio/wav",
    "audio/x-wav",
    "audio/wave",
    "audio/webm",
    "audio/ogg",
    "audio/mp4",
    "audio/m4a",
    "audio/x-m4a",
}

# Attributes that can carry user-facing UI text without it appearing as a
# text node (e.g. `<input placeholder="...">`, `<button aria-label="...">`).
_HTML_TEXT_ATTRS = ("placeholder", "aria-label", "alt", "title", "value", "label")


class SupplementalContextError(Exception):
    """Raised when supplemental context fails validation or extraction.

    `code` is a short machine-readable reason (e.g. "INVALID_BASE64",
    "UNSUPPORTED_FORMAT") that calling code can map to an HTTP response.
    """

    def __init__(self, message: str, *, code: str):
        super().__init__(message)
        self.message = message
        self.code = code


def _decode_base64(data: str) -> bytes:
    cleaned = data.strip()

    if cleaned.lower().startswith("data:") and "," in cleaned:
        # Tolerate a data: URL (e.g. "data:application/pdf;base64,JVBERi0...").
        cleaned = cleaned.split(",", 1)[1]

    try:
        return base64.b64decode(cleaned, validate=True)
    except (binascii.Error, ValueError) as exc:
        raise SupplementalContextError(
            "Supplemental content_base64 is not valid Base64.",
            code="INVALID_BASE64",
        ) from exc


def _decode_text_bytes(data: bytes) -> str:
    try:
        return data.decode("utf-8")
    except UnicodeDecodeError:
        return data.decode("latin-1", errors="replace")


def _html_to_text(html: str) -> str:
    soup = BeautifulSoup(html, "html.parser")

    for tag in soup(["script", "style", "noscript", "template"]):
        tag.decompose()

    for comment in soup.find_all(string=lambda node: isinstance(node, Comment)):
        comment.extract()

    lines: list[str] = []

    for element in soup.descendants:
        if isinstance(element, str):
            text = str(element).strip()
            if text:
                lines.append(text)
        elif getattr(element, "attrs", None):
            for attr in _HTML_TEXT_ATTRS:
                value = element.attrs.get(attr)
                if isinstance(value, str) and value.strip():
                    lines.append(value.strip())

    # De-duplicate repeated lines (common for repeated labels/buttons) while
    # preserving first-seen order.
    deduped = list(dict.fromkeys(lines))
    extracted = "\n".join(deduped)
    return re.sub(r"\n{3,}", "\n\n", extracted).strip()


def _extract_pdf_text(data: bytes) -> str:
    reader = PdfReader(io.BytesIO(data))
    pages = [(page.extract_text() or "").strip() for page in reader.pages]
    return "\n\n".join(page for page in pages if page)


def _extract_docx_text(data: bytes) -> str:
    document = Document(io.BytesIO(data))

    parts = [p.text.strip() for p in document.paragraphs if p.text.strip()]

    for table in document.tables:
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
            if cells:
                parts.append(" | ".join(cells))

    return "\n".join(parts)


def _extract_plain_text(data: bytes) -> str:
    return _decode_text_bytes(data).strip()


def _extract_json_text(data: bytes) -> str:
    text = _decode_text_bytes(data)
    try:
        parsed = json.loads(text)
    except json.JSONDecodeError as exc:
        raise SupplementalContextError(
            f"Uploaded JSON file is not valid JSON: {exc}",
            code="EXTRACTION_FAILED",
        ) from exc
    return json.dumps(parsed, indent=2, ensure_ascii=False)


def _extract_csv_text(data: bytes) -> str:
    text = _decode_text_bytes(data)
    rows = [
        row
        for row in csv.reader(io.StringIO(text))
        if any(cell.strip() for cell in row)
    ]
    return "\n".join(" | ".join(cell.strip() for cell in row) for row in rows)


def _extract_html_bytes(data: bytes) -> str:
    return _html_to_text(_decode_text_bytes(data))


FILE_EXTENSION_HANDLERS: dict[str, Callable[[bytes], str]] = {
    ".pdf": _extract_pdf_text,
    ".docx": _extract_docx_text,
    ".txt": _extract_plain_text,
    ".md": _extract_plain_text,
    ".markdown": _extract_plain_text,
    ".json": _extract_json_text,
    ".csv": _extract_csv_text,
    ".html": _extract_html_bytes,
    ".htm": _extract_html_bytes,
}

FILE_MIME_HANDLERS: dict[str, Callable[[bytes], str]] = {
    "application/pdf": _extract_pdf_text,
    "application/vnd.openxmlformats-officedocument.wordprocessingml.document": (
        _extract_docx_text
    ),
    "text/plain": _extract_plain_text,
    "text/markdown": _extract_plain_text,
    "application/json": _extract_json_text,
    "text/csv": _extract_csv_text,
    "text/html": _extract_html_bytes,
}


def _resolve_file_handler(
    file_name: Optional[str], mime_type: Optional[str]
) -> Optional[Callable[[bytes], str]]:
    if file_name:
        ext = os.path.splitext(file_name)[1].lower()
        handler = FILE_EXTENSION_HANDLERS.get(ext)
        if handler is not None:
            return handler

    if mime_type:
        return FILE_MIME_HANDLERS.get(mime_type.strip().lower())

    return None


async def _transcribe_audio(audio_bytes: bytes, file_name: Optional[str]) -> str:
    buffer = io.BytesIO(audio_bytes)
    # litellm/the OpenAI SDK read .name to sniff the audio format.
    buffer.name = file_name or "supplemental_audio.mp3"

    try:
        response = await litellm.atranscription(
            model=AUDIO_MODEL_NAME,
            file=buffer,
        )
    except Exception as exc:
        raise SupplementalContextError(
            "Audio transcription failed.",
            code="TRANSCRIPTION_FAILED",
        ) from exc

    transcript = getattr(response, "text", None)
    if transcript is None and isinstance(response, dict):
        transcript = response.get("text")

    if not transcript or not transcript.strip():
        raise SupplementalContextError(
            "Audio transcription produced no text.",
            code="TRANSCRIPTION_FAILED",
        )

    return transcript.strip()


async def _process_html(context: SupplementalContext) -> SupplementalContext:
    html_source = context.raw_html

    if not html_source and context.content_base64:
        html_source = _decode_text_bytes(_decode_base64(context.content_base64))

    if not html_source or not html_source.strip():
        raise SupplementalContextError(
            "HTML supplemental context is empty.", code="EMPTY_CONTENT"
        )

    extracted = _html_to_text(html_source)

    if not extracted:
        raise SupplementalContextError(
            "No meaningful text could be extracted from the supplied HTML.",
            code="EXTRACTION_FAILED",
        )

    logger.info("Processed HTML supplemental context (%d chars extracted)", len(extracted))

    return SupplementalContext(
        source_type="HTML",
        file_name=context.file_name,
        mime_type=context.mime_type or "text/html",
        extracted_text=extracted,
    )


async def _process_audio(context: SupplementalContext) -> SupplementalContext:
    if not context.content_base64:
        raise SupplementalContextError(
            "AUDIO supplemental context requires content_base64.",
            code="MISSING_CONTENT",
        )

    if context.mime_type and context.mime_type.strip().lower() not in SUPPORTED_AUDIO_MIME_TYPES:
        raise SupplementalContextError(
            f"Unsupported audio MIME type: {context.mime_type}",
            code="UNSUPPORTED_FORMAT",
        )

    audio_bytes = _decode_base64(context.content_base64)

    if not audio_bytes:
        raise SupplementalContextError(
            "Audio supplemental context is empty.", code="EMPTY_CONTENT"
        )

    if len(audio_bytes) > MAX_AUDIO_BYTES:
        raise SupplementalContextError(
            "Audio file exceeds the maximum allowed size.",
            code="OVERSIZED_CONTENT",
        )

    transcript = await _transcribe_audio(audio_bytes, context.file_name)

    logger.info("Processed AUDIO supplemental context (%d chars transcribed)", len(transcript))

    return SupplementalContext(
        source_type="AUDIO",
        file_name=context.file_name,
        mime_type=context.mime_type,
        transcript=transcript,
    )


async def _process_file(context: SupplementalContext) -> SupplementalContext:
    if not context.content_base64:
        raise SupplementalContextError(
            "FILE supplemental context requires content_base64.",
            code="MISSING_CONTENT",
        )

    handler = _resolve_file_handler(context.file_name, context.mime_type)

    if handler is None:
        raise SupplementalContextError(
            "Unsupported file format. Supported types: PDF, DOCX, TXT, "
            "Markdown, JSON, CSV, HTML.",
            code="UNSUPPORTED_FORMAT",
        )

    file_bytes = _decode_base64(context.content_base64)

    if not file_bytes:
        raise SupplementalContextError(
            "Uploaded file is empty.", code="EMPTY_CONTENT"
        )

    if len(file_bytes) > MAX_FILE_BYTES:
        raise SupplementalContextError(
            "Uploaded file exceeds the maximum allowed size.",
            code="OVERSIZED_CONTENT",
        )

    try:
        extracted = handler(file_bytes)
    except SupplementalContextError:
        raise
    except Exception as exc:
        raise SupplementalContextError(
            f"Failed to extract text from uploaded file: {exc}",
            code="EXTRACTION_FAILED",
        ) from exc

    if not extracted or not extracted.strip():
        raise SupplementalContextError(
            "No extractable text found in the uploaded file.",
            code="EXTRACTION_FAILED",
        )

    logger.info(
        "Processed FILE supplemental context '%s' (%d chars extracted)",
        context.file_name,
        len(extracted),
    )

    return SupplementalContext(
        source_type="FILE",
        file_name=context.file_name,
        mime_type=context.mime_type,
        extracted_text=extracted.strip(),
    )


async def process_supplemental_context(
    context: SupplementalContext,
) -> SupplementalContext:
    """Normalize raw supplemental context into LLM-safe text.

    Returns a new `SupplementalContext` with `content_base64` / `raw_html`
    cleared and `extracted_text` / `transcript` populated as applicable.
    Raises `SupplementalContextError` on any validation or extraction
    failure instead of failing silently.
    """

    if context.source_type == "NONE":
        return SupplementalContext(source_type="NONE")

    if context.source_type == "HTML":
        return await _process_html(context)

    if context.source_type == "AUDIO":
        return await _process_audio(context)

    if context.source_type == "FILE":
        return await _process_file(context)

    raise SupplementalContextError(
        f"Unsupported source_type: {context.source_type}",
        code="UNSUPPORTED_FORMAT",
    )
