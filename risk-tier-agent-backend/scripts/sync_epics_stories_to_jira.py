from __future__ import annotations

import argparse
import os
from pathlib import Path
from typing import Any

from dotenv import load_dotenv
from openpyxl import load_workbook

from app.integrations.jira_client import JiraClient, JiraConfig


def rows(path: Path, sheet_name: str) -> list[dict[str, Any]]:
    wb = load_workbook(path, read_only=True, data_only=True)
    ws = wb[sheet_name]
    it = ws.iter_rows(values_only=True)
    headers = [str(v).strip() for v in next(it)]
    return [dict(zip(headers, values)) for values in it if any(v is not None for v in values)]


def adf(text: str) -> dict[str, Any]:
    return {
        "type": "doc", "version": 1,
        "content": [{"type": "paragraph", "content": [{"type": "text", "text": text}]}],
    }


def select_value(value: Any) -> dict[str, str]:
    return {"value": str(value)}


def main() -> None:
    load_dotenv()
    parser = argparse.ArgumentParser()
    parser.add_argument("workbook")
    args = parser.parse_args()

    config = JiraConfig(
        base_url=os.environ["JIRA_BASE_URL"], email=os.environ["JIRA_EMAIL"],
        api_token=os.environ["JIRA_API_TOKEN"], project_key=os.environ["JIRA_PROJECT_KEY"],
        epic_issue_type=os.getenv("JIRA_EPIC_ISSUE_TYPE", "Epic"),
        story_issue_type=os.getenv("JIRA_STORY_ISSUE_TYPE", "Story"),
        epic_id_field=os.environ["JIRA_FIELD_EPIC_ID"],
        story_id_field=os.environ["JIRA_FIELD_STORY_ID"],
        module_id_field=os.environ["JIRA_FIELD_MODULE_ID"],
        change_status_field=os.environ["JIRA_FIELD_CHANGE_STATUS"],
        release_label=os.getenv("JIRA_RELEASE_LABEL", ""),
    )
    client = JiraClient(config)
    workbook = Path(args.workbook)

    epic_rows = rows(workbook, "02_Epics")
    story_rows = rows(workbook, "03_User_Stories")

    existing_epics = client.search(
        f'project = "{config.project_key}" AND issuetype = "{config.epic_issue_type}"',
        ["summary", config.epic_id_field],
    )
    epic_key_by_external_id = {
        str(issue.get("fields", {}).get(config.epic_id_field) or "").strip(): issue["key"]
        for issue in existing_epics
    }

    for row in epic_rows:
        external_id = str(row["Epic_ID"])
        if external_id in epic_key_by_external_id:
            print("Epic exists:", external_id, epic_key_by_external_id[external_id])
            continue
        fields = {
            "project": {"key": config.project_key},
            "issuetype": {"name": config.epic_issue_type},
            "summary": f"[{external_id}] {row['Epic_Name']}",
            "description": adf(f"This Epic covers {row['Epic_Name']}.") ,
            config.epic_id_field: external_id,
            config.module_id_field: select_value(row["Module_ID"]),
            config.change_status_field: select_value(row["Change_Status"]),
            "labels": [config.release_label] if config.release_label else [],
        }
        created = client.create_issue(fields)
        epic_key_by_external_id[external_id] = created["key"]
        print("Created Epic:", external_id, created["key"])

    existing_stories = client.search(
        f'project = "{config.project_key}" AND issuetype = "{config.story_issue_type}"',
        ["summary", config.story_id_field],
    )
    story_ids = {
        str(issue.get("fields", {}).get(config.story_id_field) or "").strip()
        for issue in existing_stories
    }

    for row in story_rows:
        story_id = str(row["Story_ID"])
        if story_id in story_ids:
            print("Story exists:", story_id)
            continue
        parent_key = epic_key_by_external_id[str(row["Epic_ID"])]
        fields = {
            "project": {"key": config.project_key},
            "issuetype": {"name": config.story_issue_type},
            "parent": {"key": parent_key},
            "summary": f"[{story_id}] {row['Story_Description']}",
            "description": adf(str(row["Story_Description"])),
            config.story_id_field: story_id,
            config.change_status_field: select_value(row["Change_Status"]),
            "labels": [config.release_label] if config.release_label else [],
        }
        created = client.create_issue(fields)
        print("Created Story:", story_id, created["key"], "parent", parent_key)


if __name__ == "__main__":
    main()
