from __future__ import annotations

from pathlib import Path
from typing import Any

from openpyxl import load_workbook


class ExcelRepository:
    """Read-only repository over the Risk Tier Agent workbook."""

    def __init__(
        self,
        workbook_path: str | Path,
    ):
        self.workbook_path = Path(
            workbook_path
        )

        if not self.workbook_path.exists():
            raise FileNotFoundError(
                "Workbook not found: "
                f"{self.workbook_path.resolve()}"
            )

        self._cache: dict[
            str,
            list[dict[str, Any]],
        ] = {}

    def rows(
        self,
        sheet_name: str,
    ) -> list[dict[str, Any]]:
        if sheet_name in self._cache:
            return self._cache[sheet_name]

        workbook = load_workbook(
            self.workbook_path,
            read_only=True,
            data_only=True,
        )

        try:
            if sheet_name not in workbook.sheetnames:
                raise KeyError(
                    f"Missing worksheet: {sheet_name}"
                )

            worksheet = workbook[sheet_name]

            values = list(
                worksheet.iter_rows(
                    values_only=True
                )
            )

            if not values:
                self._cache[sheet_name] = []
                return []

            headers = [
                str(value).strip()
                if value is not None
                else ""
                for value in values[0]
            ]

            records: list[
                dict[str, Any]
            ] = []

            for row in values[1:]:
                if all(
                    value is None
                    for value in row
                ):
                    continue

                records.append(
                    {
                        headers[index]: (
                            row[index]
                            if index < len(row)
                            else None
                        )
                        for index in range(
                            len(headers)
                        )
                    }
                )

            self._cache[sheet_name] = records
            return records

        finally:
            workbook.close()

    def filter(
        self,
        sheet_name: str,
        **criteria: Any,
    ) -> list[dict[str, Any]]:
        return [
            row
            for row in self.rows(sheet_name)
            if all(
                row.get(column) == expected
                for column, expected
                in criteria.items()
            )
        ]

    def first(
        self,
        sheet_name: str,
        **criteria: Any,
    ) -> dict[str, Any] | None:
        records = self.filter(
            sheet_name,
            **criteria,
        )

        return records[0] if records else None

    def sheet_names(self) -> list[str]:
        workbook = load_workbook(
            self.workbook_path,
            read_only=True,
        )

        try:
            return list(workbook.sheetnames)
        finally:
            workbook.close()